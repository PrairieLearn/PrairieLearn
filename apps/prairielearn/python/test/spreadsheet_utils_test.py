import datetime
from dataclasses import FrozenInstanceError, is_dataclass
from pathlib import Path
from typing import Any, cast

import numpy as np
import pandas as pd
import prairielearn as pl
import pytest
from openpyxl import Workbook
from prairielearn import spreadsheet_utils


def snapshot() -> dict[str, Any]:
    return {
        "schema_version": 1,
        "template_hash": "template-hash",
        "engine": {
            "name": "hyperformula",
            "version": "3.4.0",
            "configuration_version": 1,
        },
        "sheets": [
            {
                "name": "Inputs",
                "rows": 3,
                "columns": 3,
                "cells": {
                    "A1": {
                        "input": {"type": "number", "value": 3},
                        "result": {"type": "number", "value": 3},
                    },
                    "B1": {
                        "input": {"type": "formula", "value": "=A1*2"},
                        "result": {"type": "number", "value": 6},
                    },
                    "C1": {
                        "input": {"type": "formula", "value": "=1/0"},
                        "result": {
                            "type": "error",
                            "value": "#DIV/0!",
                            "error_type": "DIV_BY_ZERO",
                        },
                    },
                },
            },
            {
                "name": "Input Data",
                "rows": 2,
                "columns": 2,
                "cells": {
                    "A1": {
                        "input": {"type": "number", "value": 10},
                        "result": {"type": "number", "value": 10},
                    },
                    "B2": {
                        "input": {"type": "formula", "value": "=A1+5"},
                        "result": {"type": "number", "value": 15},
                    },
                },
            },
            {
                "name": "Bob's Data",
                "rows": 1,
                "columns": 1,
                "cells": {
                    "A1": {
                        "input": {"type": "string", "value": "quoted"},
                        "result": {"type": "string", "value": "quoted"},
                    }
                },
            },
        ],
        "grading": {
            "schema_version": 1,
            "grader_hash": "grader-hash",
            "outputs": {
                "total": {"type": "number", "value": 6},
                "failed_check": {
                    "type": "error",
                    "value": "#N/A",
                    "error_type": "NA",
                },
            },
        },
    }


@pytest.mark.parametrize(
    ("address", "expected"),
    [
        ("a1", ("A1", 0, 0)),
        ("aa10", ("AA10", 9, 26)),
        ("zZ99", ("ZZ99", 98, 701)),
    ],
)
def test_normalized_address_accepts_case_insensitive_columns(
    address: str, expected: tuple[str, int, int]
) -> None:
    assert spreadsheet_utils._normalized_address(address) == expected


def test_ranges_report_shape_and_containment() -> None:
    local_range = spreadsheet_utils.SpreadsheetAddressRange(
        spreadsheet_utils.SpreadsheetAddress(row=1, column=2),
        spreadsheet_utils.SpreadsheetAddress(row=3, column=5),
    )

    assert local_range.shape == (3, 4)
    assert spreadsheet_utils.SpreadsheetAddress(row=1, column=2) in local_range
    assert spreadsheet_utils.SpreadsheetAddress(row=3, column=5) in local_range
    assert spreadsheet_utils.SpreadsheetAddress(row=0, column=2) not in local_range
    assert spreadsheet_utils.SpreadsheetAddress(row=3, column=6) not in local_range
    assert (
        spreadsheet_utils.SpreadsheetAddressRange(
            spreadsheet_utils.SpreadsheetAddress(row=2, column=3),
            spreadsheet_utils.SpreadsheetAddress(row=3, column=5),
        )
        in local_range
    )

    qualified_range = spreadsheet_utils.SpreadsheetQualifiedAddressRange(
        local_range, "Inputs"
    )
    assert qualified_range.shape == (3, 4)
    assert (
        spreadsheet_utils.SpreadsheetQualifiedAddress(
            spreadsheet_utils.SpreadsheetAddress(row=2, column=4), "Inputs"
        )
        in qualified_range
    )
    assert spreadsheet_utils.SpreadsheetAddress(row=2, column=4) in qualified_range
    assert (
        spreadsheet_utils.SpreadsheetQualifiedAddress(
            spreadsheet_utils.SpreadsheetAddress(row=2, column=4), "Other"
        )
        not in qualified_range
    )
    assert (
        spreadsheet_utils.SpreadsheetAddressRange(
            spreadsheet_utils.SpreadsheetAddress(row=2, column=3),
            spreadsheet_utils.SpreadsheetAddress(row=3, column=5),
        )
        in qualified_range
    )
    assert (
        spreadsheet_utils.SpreadsheetQualifiedAddressRange(local_range, "Inputs")
        in qualified_range
    )
    assert (
        spreadsheet_utils.SpreadsheetQualifiedAddressRange(local_range, "Other")
        not in qualified_range
    )


def test_range_intersections() -> None:
    local_range = spreadsheet_utils.SpreadsheetAddressRange(
        spreadsheet_utils.SpreadsheetAddress(row=1, column=2),
        spreadsheet_utils.SpreadsheetAddress(row=4, column=5),
    )
    overlapping = spreadsheet_utils.SpreadsheetAddressRange(
        spreadsheet_utils.SpreadsheetAddress(row=3, column=1),
        spreadsheet_utils.SpreadsheetAddress(row=5, column=3),
    )
    expected_local = spreadsheet_utils.SpreadsheetAddressRange(
        spreadsheet_utils.SpreadsheetAddress(row=3, column=2),
        spreadsheet_utils.SpreadsheetAddress(row=4, column=3),
    )
    disjoint = spreadsheet_utils.SpreadsheetAddressRange(
        spreadsheet_utils.SpreadsheetAddress(row=5, column=2),
        spreadsheet_utils.SpreadsheetAddress(row=6, column=5),
    )

    assert local_range.intersection(overlapping) == expected_local
    assert local_range.intersection(disjoint) is None

    qualified_range = spreadsheet_utils.SpreadsheetQualifiedAddressRange(
        local_range, "Inputs"
    )
    expected_qualified = spreadsheet_utils.SpreadsheetQualifiedAddressRange(
        expected_local, "Inputs"
    )
    assert qualified_range.intersection(overlapping) == expected_qualified
    assert (
        qualified_range.intersection(
            spreadsheet_utils.SpreadsheetQualifiedAddressRange(overlapping, "Inputs")
        )
        == expected_qualified
    )
    assert (
        qualified_range.intersection(
            spreadsheet_utils.SpreadsheetQualifiedAddressRange(overlapping, "Other")
        )
        is None
    )
    assert (
        local_range.intersection(
            spreadsheet_utils.SpreadsheetQualifiedAddressRange(overlapping, "Other")
        )
        == expected_local
    )


def test_address_ranges_parse_and_translate_relative_coordinates() -> None:
    cell_range = pl.SpreadsheetAddressRange.from_a1("D5:B3")

    assert cell_range.address == "B3:D5"
    assert cell_range.shape == (3, 3)
    assert cell_range.to_source(pl.SpreadsheetAddress(1, 2)).address == "D4"
    assert cell_range.to_relative(pl.SpreadsheetAddress.from_a1("C5")) == (
        pl.SpreadsheetAddress(2, 1)
    )
    with pytest.raises(ValueError, match="outside"):
        cell_range.to_source(pl.SpreadsheetAddress(3, 0))


def test_dataframe_conversion_is_sparse_json_safe_and_offset() -> None:
    dataframe = pd.DataFrame(
        [[np.int64(2), "=C3*2"], [pd.NA, np.float64(3.5)]],
        columns=["quantity", "formula"],
        index=["first", "second"],
    )

    sheet = pl.dataframe_to_spreadsheet_sheet(dataframe, name="Inputs", start_cell="C3")
    assert sheet == {
        "name": "Inputs",
        "rows": 4,
        "columns": 4,
        "cells": {"C3": 2, "D3": "=C3*2", "D4": 3.5},
        "editable_ranges": [],
    }

    labeled = pl.dataframe_to_spreadsheet_sheet(
        dataframe,
        start_cell="B2",
        include_columns=True,
        include_index=True,
        editable_ranges=("C3:D4",),
    )
    assert labeled["cells"] == {
        "C2": "quantity",
        "D2": "formula",
        "B3": "first",
        "B4": "second",
        "C3": 2,
        "D3": "=C3*2",
        "D4": 3.5,
    }
    assert labeled.get("editable_ranges") == ["C3:D4"]


@pytest.mark.parametrize(
    "value", [float("inf"), 1 + 2j, {"nested": True}, datetime.date(2024, 1, 1)]
)
def test_dataframe_conversion_rejects_non_json_cell_values(value: object) -> None:
    with pytest.raises((TypeError, ValueError), match="Spreadsheet value"):
        pl.dataframe_to_spreadsheet_sheet(pd.DataFrame([[value]]))


def test_dataframe_conversion_rejects_requested_multi_index_labels() -> None:
    columns = pd.DataFrame(
        [[1]], columns=pd.MultiIndex.from_tuples([("group", "value")])
    )
    index = pd.DataFrame([[1]], index=pd.MultiIndex.from_tuples([("group", "row")]))

    with pytest.raises(TypeError, match="column labels"):
        pl.dataframe_to_spreadsheet_sheet(columns, include_columns=True)
    with pytest.raises(TypeError, match="index labels"):
        pl.dataframe_to_spreadsheet_sheet(index, include_index=True)


def test_dataframe_book_rejects_case_insensitive_sheet_collisions() -> None:
    frames = {"Inputs": pd.DataFrame([[1]]), "inputs": pd.DataFrame([[2]])}

    with pytest.raises(ValueError, match="duplicated"):
        pl.dataframes_to_spreadsheet_book(frames)


def test_public_file_readers_load_csv_tsv_and_xlsx(tmp_path: Path) -> None:
    csv_path = tmp_path / "input.csv"
    csv_path.write_text("2,=A1*2\n,NA\n")
    tsv_path = tmp_path / "input.tsv"
    tsv_path.write_text("2\t=A1*2\n\tNA\n")

    for source, reader in (
        (csv_path, pl.read_spreadsheet_csv),
        (tsv_path, pl.read_spreadsheet_tsv),
    ):
        book = reader(source, sheet_name="Inputs")
        assert book["sheets"][0]["cells"] == {
            "A1": "2",
            "B1": "=A1*2",
            "B2": "NA",
        }

    xlsx_path = tmp_path / "input.xlsx"
    workbook = Workbook()
    inputs = workbook.active
    assert inputs is not None
    inputs.title = "Inputs"
    inputs["A1"] = 2
    inputs["B1"] = "=A1*2"
    checks = workbook.create_sheet("Checks")
    checks["A1"] = "=Inputs!B1=4"
    workbook.save(xlsx_path)

    book = pl.read_spreadsheet(xlsx_path)
    assert [sheet["name"] for sheet in book["sheets"]] == ["Inputs", "Checks"]
    assert book["sheets"][0]["cells"]["B1"] == "=A1*2"
    assert book["sheets"][1]["cells"]["A1"] == "=Inputs!B1=4"


def test_sparse_cells_data_contains_valid_cells() -> None:
    cells = spreadsheet_utils._SparseCellsData(
        name="Input Data",
        rows=3,
        columns=4,
        visible_range=pl.SpreadsheetAddressRange.from_a1("A1:D3"),
        addressed_data={},
    )

    assert spreadsheet_utils.SpreadsheetAddress(row=0, column=0) in cells
    assert spreadsheet_utils.SpreadsheetAddress(row=2, column=3) in cells
    assert spreadsheet_utils.SpreadsheetAddress(row=3, column=3) not in cells
    assert (
        spreadsheet_utils.SpreadsheetQualifiedAddress(
            spreadsheet_utils.SpreadsheetAddress(row=2, column=3), "Input Data"
        )
        in cells
    )
    assert (
        spreadsheet_utils.SpreadsheetQualifiedAddress(
            spreadsheet_utils.SpreadsheetAddress(row=2, column=3), "Other"
        )
        not in cells
    )


def test_queries_cells_formulas_and_outputs() -> None:
    workbook = snapshot()

    assert pl.get_spreadsheet_cell(workbook, "Inputs", "a1") == {
        "input": {"type": "number", "value": 3},
        "result": {"type": "number", "value": 3},
    }
    assert pl.get_spreadsheet_cell(workbook, "Inputs", "A2") is None
    assert pl.get_spreadsheet_result(workbook, "Inputs", "A2") == {"type": "empty"}
    assert pl.get_spreadsheet_value(workbook, "Inputs", "B1") == 6
    assert pl.get_spreadsheet_value(workbook, "Inputs", "C1") is None
    assert pl.get_spreadsheet_formula(workbook, "Inputs", "B1") == "=A1*2"
    assert pl.get_spreadsheet_formula(workbook, "Inputs", "A1") is None
    assert pl.get_spreadsheet_grading_output(workbook, "total") == {
        "type": "number",
        "value": 6,
    }


def test_query_errors_are_explicit() -> None:
    workbook = snapshot()

    with pytest.raises(KeyError, match="Unknown spreadsheet sheet"):
        pl.get_spreadsheet_cell(workbook, "Missing", "A1")
    with pytest.raises(ValueError, match="outside the sheet"):
        pl.get_spreadsheet_cell(workbook, "Inputs", "D1")
    with pytest.raises(ValueError, match="Invalid spreadsheet cell address"):
        pl.get_spreadsheet_cell(workbook, "Inputs", "not-a-cell")
    with pytest.raises(KeyError, match="Unknown spreadsheet grading output"):
        pl.get_spreadsheet_grading_output(workbook, "missing")
    workbook.pop("grading")
    with pytest.raises(KeyError, match="does not contain private grading outputs"):
        pl.get_spreadsheet_grading_output(workbook, "total")


def test_spreadsheet_wrapper_addresses_cells_and_sheets() -> None:
    workbook = pl.SpreadsheetBook(snapshot())

    assert workbook.sheet_names == ("Inputs", "Input Data", "Bob's Data")
    inputs = workbook["Inputs"]
    assert isinstance(inputs, pl.Spreadsheet)
    assert inputs.book is workbook
    assert inputs.name == "Inputs"
    assert inputs.shape == (3, 3)
    assert inputs.rows == 3
    assert inputs.columns == 3

    cell = inputs.cell("a1")
    assert cell.sheet is inputs
    assert cell.address == "A1"
    assert cell.qualified_address == "Inputs!A1"
    assert cell.row == 1
    assert cell.column == "A"
    assert cell.input == {"type": "number", "value": 3}
    assert cell.result == {"type": "number", "value": 3}
    assert cell.value == 3
    assert not cell.is_empty
    assert not cell.is_formula
    assert not cell.is_error
    assert cell.error_type is None
    assert cell.error_value is None

    empty = inputs["A2"]
    assert isinstance(empty, pl.SpreadsheetCellView)
    assert empty.input is None
    assert empty.result == {"type": "empty"}
    assert empty.value is None
    assert empty.is_empty
    assert not empty.is_error
    assert empty.error_type is None
    assert empty.error_value is None

    error = inputs.cell("C1")
    assert error.value is None
    assert not error.is_empty
    assert error.is_error
    assert error.error_type == "DIV_BY_ZERO"
    assert error.error_value == "#DIV/0!"


def test_spreadsheet_wrapper_constrains_offset_student_range() -> None:
    offset_snapshot = snapshot()
    sheet = offset_snapshot["sheets"][0]
    sheet["student_range"] = "B1:C2"
    sheet["cells"].pop("A1")

    inputs = pl.SpreadsheetBook(offset_snapshot)["Inputs"]
    assert inputs.shape == (2, 2)
    assert tuple(cell.address for cell in inputs) == ("B1", "C1", "B2", "C2")
    assert tuple(
        cell.address for cell in inputs.query(lambda cell: cell.is_formula)
    ) == (
        "B1",
        "C1",
    )
    assert inputs.cell("B1").value == 6
    with pytest.raises(ValueError, match="outside the sheet"):
        inputs["A1"]
    with pytest.raises(ValueError, match="outside the sheet"):
        pl.get_spreadsheet_cell(offset_snapshot, "Inputs", "A1")

    forged_snapshot = snapshot()
    forged_snapshot["sheets"][0]["student_range"] = "B1:C2"
    with pytest.raises(ValueError, match="outside the student range"):
        pl.SpreadsheetBook(forged_snapshot)


def test_spreadsheet_validates_and_copies_snapshot_at_construction() -> None:
    invalid_input = snapshot()
    invalid_input["sheets"][0]["cells"]["A1"]["input"] = {
        "type": "number",
        "value": "not a number",
    }
    with pytest.raises(TypeError, match="invalid input"):
        pl.SpreadsheetBook(invalid_input)

    invalid_output = snapshot()
    invalid_output["grading"]["outputs"]["total"] = {
        "type": "number",
        "value": "not a number",
    }
    with pytest.raises(TypeError, match="invalid result"):
        pl.SpreadsheetBook(invalid_output)

    original = snapshot()
    workbook = pl.SpreadsheetBook(original)
    original["sheets"][0]["cells"]["A1"]["result"]["value"] = 99
    assert workbook["Inputs"].cell("A1").value == 3


def test_spreadsheet_wrapper_views_are_frozen_slots_dataclasses() -> None:
    workbook = pl.SpreadsheetBook(snapshot())
    sheet = workbook["Inputs"]
    cell = sheet["A1"]
    cell_range = sheet["A1:B2"]
    assert isinstance(cell, pl.SpreadsheetCellView)
    assert isinstance(cell_range, pl.SpreadsheetRange)

    views = (
        workbook,
        sheet,
        cell,
        cell_range,
        workbook.outputs,
        workbook.outputs["total"],
    )
    assert all(is_dataclass(view) for view in views)
    assert all(not hasattr(view, "__dict__") for view in views)

    set_attribute = setattr
    with pytest.raises(FrozenInstanceError):
        set_attribute(sheet, "book", workbook)


def test_spreadsheet_wrapper_resolves_cross_sheet_references() -> None:
    workbook = pl.SpreadsheetBook(snapshot())
    inputs = workbook["Inputs"]

    other_range = inputs.range("'Input Data'!A1:B2")
    assert other_range.sheet is workbook["Input Data"]
    assert other_range.range_address == "A1:B2"
    assert other_range.qualified_range_address == "'Input Data'!A1:B2"
    assert other_range.shape == (2, 2)
    assert not hasattr(other_range, "range")
    assert not hasattr(other_range, "start")
    assert not hasattr(other_range, "end")
    assert other_range.values == ((10, None), (None, 15))  # ruff: ignore[pandas-use-of-dot-values]
    assert other_range["B2"].value == 15
    assert other_range["'Input Data'!A1"].value == 10

    quoted_cell = inputs["'Bob''s Data'!A1"]
    assert isinstance(quoted_cell, pl.SpreadsheetCellView)
    assert quoted_cell.value == "quoted"
    assert quoted_cell.qualified_address == "'Bob''s Data'!A1"

    with pytest.raises(TypeError, match="is not a cell"):
        inputs.cell("A1:B2")
    with pytest.raises(TypeError, match="is not a range"):
        inputs.range("A1")


@pytest.mark.parametrize(
    ("reference", "error"),
    [
        ("Inputs!A1:'Input Data'!B2", "cannot span multiple sheets"),
        ("A1:'Input Data'!B2", "cannot span multiple sheets"),
        ("B2:A1", "top-left to bottom-right"),
        ("'Input Data!A1", "Invalid spreadsheet reference"),
        ("A:A", "Invalid spreadsheet reference"),
        ("A1:B2:C3", "Invalid spreadsheet reference"),
    ],
)
def test_spreadsheet_wrapper_rejects_invalid_references(
    reference: str, error: str
) -> None:
    inputs = pl.SpreadsheetBook(snapshot())["Inputs"]

    with pytest.raises(ValueError, match=error):
        inputs[reference]

    with pytest.raises(KeyError, match="Unknown spreadsheet sheet"):
        inputs["Missing!A1"]
    with pytest.raises(ValueError, match="outside the sheet"):
        inputs["D1"]


def test_spreadsheet_range_projections_queries_and_indexing() -> None:
    inputs = pl.SpreadsheetBook(snapshot())["Inputs"]
    cell_range = inputs["A1:B2"]
    assert isinstance(cell_range, pl.SpreadsheetRange)

    assert tuple(cell.address for cell in cell_range.iter_cells()) == (
        "A1",
        "B1",
        "A2",
        "B2",
    )
    assert tuple(
        cell.address for cell in cell_range.iter_cells(include_empty=False)
    ) == (
        "A1",
        "B1",
    )
    assert cell_range.inputs == (
        (
            {"type": "number", "value": 3},
            {"type": "formula", "value": "=A1*2"},
        ),
        (None, None),
    )
    assert cell_range.results == (
        ({"type": "number", "value": 3}, {"type": "number", "value": 6}),
        ({"type": "empty"}, {"type": "empty"}),
    )
    assert cell_range.values == ((3, 6), (None, None))  # ruff: ignore[pandas-use-of-dot-values]
    assert inputs.range("B1:C2").values == ((6, None), (None, None))  # ruff: ignore[pandas-use-of-dot-values]
    assert cell_range.formulas == ((None, "=A1*2"), (None, None))
    assert tuple(
        cell.address for cell in cell_range.query(lambda cell: cell.is_formula)
    ) == ("B1",)
    assert tuple(
        cell.address
        for cell in cell_range.query(lambda cell: cell.is_empty, include_empty=True)
    ) == ("A2", "B2")
    assert cell_range["A1"].value == 3

    with pytest.raises(ValueError, match="outside range"):
        cell_range["C1"]
    with pytest.raises(ValueError, match="single cell"):
        cell_range["A1:B1"]

    formulas = inputs.query(lambda cell: cell.is_formula)
    assert tuple(cell.address for cell in formulas) == ("B1", "C1")


def test_spreadsheet_cell_formula_matching() -> None:
    inputs = pl.SpreadsheetBook(snapshot())["Inputs"]
    formula = inputs["B1"]
    assert isinstance(formula, pl.SpreadsheetCellView)

    assert formula.formula == "=A1*2"
    assert formula.formula_ast == pl.parse_spreadsheet_formula("=A1*2")
    assert formula.matches_formula("=A1*2")
    assert not formula.matches_formula("=A1 * 2")
    assert formula.matches_formula("=A1 * 2", structural=True)
    assert not cast(pl.SpreadsheetCellView, inputs["A1"]).matches_formula("=A1")

    with pytest.raises(pl.SpreadsheetFormulaParseError):
        formula.matches_formula("not-a-formula", structural=True)

    invalid_snapshot = snapshot()
    invalid_snapshot["sheets"][0]["cells"]["B1"] = {
        "input": {"type": "formula", "value": "=named_expression"},
        "result": {"type": "error", "value": "#NAME?", "error_type": "NAME"},
    }
    invalid_formula = pl.SpreadsheetBook(invalid_snapshot)["Inputs"].cell("B1")
    assert not invalid_formula.matches_formula("=A1*2", structural=True)


def test_spreadsheet_output_views() -> None:
    workbook = pl.SpreadsheetBook(snapshot())

    assert tuple(workbook.outputs) == ("total", "failed_check")
    total = workbook.outputs["total"]
    assert total.name == "total"
    assert total.result == {"type": "number", "value": 6}
    assert total.value == 6
    assert not total.is_empty
    assert not total.is_error
    assert total.error_type is None
    assert total.error_value is None

    failed = workbook.outputs["failed_check"]
    assert failed.value is None
    assert not failed.is_empty
    assert failed.is_error
    assert failed.error_type == "NA"
    assert failed.error_value == "#N/A"

    empty_snapshot = snapshot()
    empty_snapshot["grading"]["outputs"]["empty_check"] = {"type": "empty"}
    empty = pl.SpreadsheetBook(empty_snapshot).outputs["empty_check"]
    assert empty.value is None
    assert empty.is_empty
    assert not empty.is_error
    assert empty.error_type is None
    assert empty.error_value is None

    without_outputs = snapshot()
    without_outputs.pop("grading")
    outputs = pl.SpreadsheetBook(without_outputs).outputs
    assert len(outputs) == 0
    with pytest.raises(KeyError, match="does not contain private grading outputs"):
        outputs["total"]


def test_formula_ast_respects_operator_precedence() -> None:
    ast = pl.parse_spreadsheet_formula("=A1+B2*C3^2")
    root = ast.root

    assert ast.schema_version == 1
    assert ast.formula == "=A1+B2*C3^2"
    assert isinstance(root, pl.FormulaBinaryNode)
    assert root.type == "binary"
    assert root.operator == "+"
    assert isinstance(root.right, pl.FormulaBinaryNode)
    assert root.right.operator == "*"
    assert isinstance(root.right.right, pl.FormulaBinaryNode)
    assert root.right.right.operator == "^"
    assert not hasattr(ast, "__dict__")
    assert not hasattr(root, "__dict__")
    set_attribute = setattr
    with pytest.raises(FrozenInstanceError):
        set_attribute(ast, "formula", "=1")


def test_formula_ast_preserves_formula_structure_and_references() -> None:
    ast = pl.parse_spreadsheet_formula("=IF('Sheet Name'!$A1:B$2>=3,\"ok\",)")
    root = ast.root

    assert isinstance(root, pl.FormulaFunctionNode)
    assert root.type == "function"
    assert root.name == "IF"
    assert len(root.arguments) == 3
    comparison = root.arguments[0]
    assert isinstance(comparison, pl.FormulaBinaryNode)
    assert comparison.operator == ">="
    reference_range = comparison.left
    assert isinstance(reference_range, pl.FormulaCellRangeNode)
    assert reference_range.start.kind == "cell"
    assert reference_range.start.row == 1
    assert reference_range.end.kind == "cell"
    assert reference_range.end.column == "B"
    assert reference_range == pl.FormulaCellRangeNode(
        start=pl.FormulaCellReference(
            sheet="Sheet Name",
            column="A",
            row=1,
            column_absolute=True,
            row_absolute=False,
        ),
        end=pl.FormulaCellReference(
            sheet=None,
            column="B",
            row=2,
            column_absolute=False,
            row_absolute=True,
        ),
    )
    assert isinstance(root.arguments[2], pl.FormulaEmptyNode)


def test_formula_ast_handles_unary_postfix_and_full_ranges() -> None:
    unary = pl.parse_spreadsheet_formula("=-A1%")
    assert isinstance(unary.root, pl.FormulaUnaryNode)
    assert isinstance(unary.root.operand, pl.FormulaPostfixNode)

    ranges = pl.parse_spreadsheet_formula("=SUM($A:$B,1:2)")
    assert isinstance(ranges.root, pl.FormulaFunctionNode)
    column_range = ranges.root.arguments[0]
    row_range = ranges.root.arguments[1]
    assert isinstance(column_range, pl.FormulaColumnRangeNode)
    assert isinstance(row_range, pl.FormulaRowRangeNode)
    assert column_range.start.kind == "column"
    assert column_range.start.column == "A"
    assert column_range.start.column_absolute is True
    assert column_range.end.kind == "column"
    assert column_range.end.column == "B"
    assert row_range.start.kind == "row"
    assert row_range.start.row == 1
    assert row_range.end.kind == "row"
    assert row_range.end.row == 2


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        (
            '=(1+2)&"x"',
            pl.FormulaBinaryNode(
                operator="&",
                left=pl.FormulaGroupNode(
                    expression=pl.FormulaBinaryNode(
                        operator="+",
                        left=pl.FormulaLiteralNode(value_type="number", value=1),
                        right=pl.FormulaLiteralNode(value_type="number", value=2),
                    )
                ),
                right=pl.FormulaLiteralNode(value_type="string", value="x"),
            ),
        ),
        (
            "=TRUE=FALSE",
            pl.FormulaBinaryNode(
                operator="=",
                left=pl.FormulaLiteralNode(value_type="boolean", value=True),
                right=pl.FormulaLiteralNode(value_type="boolean", value=False),
            ),
        ),
        (
            "=#N/A",
            pl.FormulaLiteralNode(value_type="error", value="#N/A"),
        ),
        (
            "=sum(A1)",
            pl.FormulaFunctionNode(
                name="SUM",
                arguments=(
                    pl.FormulaReferenceNode(
                        reference=pl.FormulaCellReference(
                            sheet=None,
                            column="A",
                            row=1,
                            column_absolute=False,
                            row_absolute=False,
                        )
                    ),
                ),
            ),
        ),
    ],
)
def test_formula_ast_golden_nodes(formula: str, expected: pl.FormulaAstNode) -> None:
    assert pl.parse_spreadsheet_formula(formula).root == expected


def test_get_formula_ast_and_invalid_formula() -> None:
    ast = pl.get_spreadsheet_formula_ast(snapshot(), "Inputs", "B1")
    assert ast is not None
    assert ast.formula == "=A1*2"
    assert pl.get_spreadsheet_formula_ast(snapshot(), "Inputs", "A1") is None

    with pytest.raises(pl.SpreadsheetFormulaParseError, match="must start"):
        pl.parse_spreadsheet_formula("A1+1")
    with pytest.raises(pl.SpreadsheetFormulaParseError, match="Unsupported"):
        pl.parse_spreadsheet_formula("=named_expression")
    with pytest.raises(pl.SpreadsheetFormulaParseError, match="incompatible endpoint"):
        pl.parse_spreadsheet_formula("=A1:B")
